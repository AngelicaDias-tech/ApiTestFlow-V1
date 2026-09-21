from pathlib import Path
from typing import BinaryIO

import openpyxl
from fastapi import HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from app.core.config import BASE_DIR
from app.models.massa import MassaImportacao
from app.repositories import api_test_repository, massa_repository

MODELO_MASSA_PATH = BASE_DIR.parent / "modelo_atualizacao_request.xlsx"
ABA_ATUALIZACOES = "Atualizacoes"
COLUNAS_ESPERADAS = ["caso_id", "path", "acao", "tipo", "novo_valor", "ativo", "observacao"]


def caminho_modelo() -> Path:
    if not MODELO_MASSA_PATH.exists():
        raise HTTPException(
            status.HTTP_404_NOT_FOUND,
            detail="Modelo oficial de massa (modelo_atualizacao_request.xlsx) não encontrado no servidor.",
        )
    return MODELO_MASSA_PATH


def _garantir_teste_existe(db: Session, api_test_id: int) -> None:
    if api_test_repository.obter(db, api_test_id) is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Teste não encontrado.")


def obter_por_teste(db: Session, api_test_id: int) -> MassaImportacao | None:
    _garantir_teste_existe(db, api_test_id)
    return massa_repository.obter_por_teste(db, api_test_id)


def _linha_valida(valores: dict) -> bool:
    return bool(valores.get("caso_id")) and bool(valores.get("path"))


def _parse_ativo(valor) -> bool:
    if isinstance(valor, bool):
        return valor
    texto = str(valor or "").strip().lower()
    return texto in ("sim", "true", "1", "yes")


def _ler_planilha(arquivo: BinaryIO) -> list[dict]:
    try:
        workbook = openpyxl.load_workbook(arquivo, data_only=True)
    except Exception as exc:  # arquivo corrompido ou não é um .xlsx válido
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Não foi possível ler o arquivo enviado como planilha .xlsx."
        ) from exc

    if ABA_ATUALIZACOES not in workbook.sheetnames:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=f"A planilha precisa conter a aba '{ABA_ATUALIZACOES}' do modelo oficial.",
        )

    aba = workbook[ABA_ATUALIZACOES]
    linhas_iter = aba.iter_rows(values_only=True)
    cabecalho = next(linhas_iter, None)
    if cabecalho is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, detail="A aba de atualizações está vazia.")

    cabecalho_normalizado = [str(c).strip() if c is not None else "" for c in cabecalho]
    if cabecalho_normalizado[: len(COLUNAS_ESPERADAS)] != COLUNAS_ESPERADAS:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=(
                "Cabeçalho da planilha não confere com o modelo oficial. Esperado: "
                + ", ".join(COLUNAS_ESPERADAS)
            ),
        )

    linhas: list[dict] = []
    for row in linhas_iter:
        if row is None or all(valor is None for valor in row):
            continue
        valores = dict(zip(COLUNAS_ESPERADAS, row))
        registro = {
            "caso_id": str(valores.get("caso_id") or "").strip(),
            "path": str(valores.get("path") or "").strip(),
            "acao": str(valores.get("acao") or "set").strip().lower(),
            "tipo": (str(valores["tipo"]).strip() if valores.get("tipo") is not None else None),
            "novo_valor": (str(valores["novo_valor"]) if valores.get("novo_valor") is not None else None),
            "ativo": _parse_ativo(valores.get("ativo")),
            "observacao": (str(valores["observacao"]) if valores.get("observacao") is not None else None),
        }
        if _linha_valida(registro):
            linhas.append(registro)

    if not linhas:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Nenhuma linha válida (com caso_id e path) foi encontrada."
        )

    return linhas


def importar(db: Session, api_test_id: int, arquivo: UploadFile) -> MassaImportacao:
    _garantir_teste_existe(db, api_test_id)

    if not arquivo.filename or not arquivo.filename.lower().endswith(".xlsx"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Envie um arquivo .xlsx.")

    linhas = _ler_planilha(arquivo.file)
    return massa_repository.substituir(db, api_test_id, arquivo.filename, linhas)
